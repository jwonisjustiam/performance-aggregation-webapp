from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from processors.samsung_processor import process_samsung
from processors.weekly_processor import process_detail, process_weekly
from services.excel_writer import create_result_workbook
from services.result_formatter import build_download_filename, shorten_model
from services.validator import validate_saved_workbook


def test_model_shortening() -> None:
    assert shorten_model("SM-R390NZSAKOO") == "SM-R390"
    assert shorten_model("SM-L705NAW1KOO") == "SM-L705"


def test_weekly_download_filename_includes_type_and_date() -> None:
    result = {"final": pd.DataFrame({"날짜": ["2026-06-24"]})}
    name = build_download_filename("weekly", result, pd.Series(dtype=object), "external")
    assert name == "외장하드 20260624 정리본.xlsx"


def test_weekly_download_filename_uses_date_range() -> None:
    result = {"final": pd.DataFrame({"날짜": ["2026-06-24", "2026-06-25"]})}
    name = build_download_filename("weekly", result, pd.Series(dtype=object), "wearable")
    assert name == "웨어러블 20260624-20260625 정리본.xlsx"


def test_weekly_workbook_reopens(tmp_path: Path, weekly_frame: pd.DataFrame) -> None:
    content, validation = create_result_workbook("weekly", process_weekly(weekly_frame, "외장하드.xlsx"))
    path = tmp_path / "weekly result.xlsx"
    path.write_bytes(content)
    assert validation["valid"]
    assert validate_saved_workbook(path, ["회차별 합계"])["valid"]
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["회차별 합계"]
        assert sheet["A1"].value == "방송 정보"
        assert sheet["R1"].value == "목표"
        assert sheet["U1"].value == "실적"
        assert sheet["AA1"].value == "AI라이브 대응"
        assert [sheet.cell(2, column).value for column in range(18, 25)] == [
            "View(만)", "수량", "금액(백만)", "View(만)", "수량", "전환율", "금액(백만)",
        ]
        assert sheet["R3"].value == pytest.approx(0.1)
        assert sheet["S3"].value == pytest.approx(50)
        assert sheet["T3"].value == pytest.approx(10)
        assert sheet["Z3"].value is None
        assert sheet["AA3"].value == "쇼마젠시"
        for row in range(3, sheet.max_row + 1):
            assert sheet.cell(row, 23).value == f"=V{row}/(U{row}*10000)"
            assert sheet.cell(row, 23).number_format == "0.00%"
            assert sheet.cell(row, 26).value is None
        assert sheet.freeze_panes == "A3"
        assert workbook.calculation.calcMode == "auto"
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
    finally:
        workbook.close()


def test_weekly_conversion_excel_format(weekly_frame: pd.DataFrame) -> None:
    stats = pd.DataFrame(
        [
            {
                "계정": "삼성공식파트너 쇼마젠시",
                "방송 ID": "100",
                "방송 제목": "외장하드 라이브",
                "방송일시": pd.Timestamp("2026-06-22 11:48"),
                "라이브중 시청수": 200,
                "유니크 결제자수": 5,
                "결제 상품수": 7,
                "데이터 업데이트 시각": pd.Timestamp("2026-06-23 10:00"),
            }
        ]
    )
    content, _ = create_result_workbook(
        "weekly", process_weekly(weekly_frame, "외장하드.xlsx", live_stats=stats)
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        sheet = workbook["회차별 합계"]
        headers = {cell.value: cell.column for cell in sheet[2]}
        row = next(
            row_number
            for row_number in range(3, sheet.max_row + 1)
            if sheet.cell(row_number, headers["시작 시간"]).value == "11:50"
        )
        cell = sheet.cell(row, headers["전환율"])
        assert cell.value == f"=V{row}/(U{row}*10000)"
        assert cell.number_format == "0.00%"
    finally:
        workbook.close()


def test_samsung_workbook_reopens(tmp_path: Path, samsung_frame: pd.DataFrame) -> None:
    content, validation = create_result_workbook("samsung", process_samsung(samsung_frame))
    path = tmp_path / "samsung result.xlsx"
    path.write_bytes(content)
    assert validation["valid"]
    assert validate_saved_workbook(path, ["통합 실적표", "회차별 합계", "중복 주문 검증"])["valid"]


def test_detail_workbook_reopens(tmp_path: Path) -> None:
    raw = pd.DataFrame(
        [
            ["A", "2026-07-07 14:00", "갤럭시 워치9", "", "SM-L340NZEAKOO", 1_000_000, 0],
            ["B", "2026-07-07 15:00", "충전 어댑터", "EF-QF976CTEGKR", "SELLER-B", 500_000, 0],
        ],
        columns=["주문번호", "결제일시", "상품명", "옵션관리코드", "판매자 상품코드", "상품가격", "옵션가격"],
    )
    content, validation = create_result_workbook("detail", process_detail(raw, "주문.xlsx"))
    path = tmp_path / "detail result.xlsx"
    path.write_bytes(content)
    assert validation["valid"]
    assert validate_saved_workbook(path, ["Basic", "웨어러블", "모바일 ACC"])["valid"]
