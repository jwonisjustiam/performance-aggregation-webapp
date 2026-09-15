from datetime import date, time

import pandas as pd
import pytest

from processors.weekly_processor import infer_target_dates, infer_weekly_kind, process_detail, process_weekly
from rules.weekly_rules import SlotRule
from services.excel_reader import canonicalize_columns


def test_unique_order_count_and_multiple_line_amount(weekly_frame: pd.DataFrame) -> None:
    result = process_weekly(weekly_frame, "외장하드 주문.xlsx")
    row = result["final"].query("`시작 시간` == '11:50'").iloc[0]
    assert row["실적 수량"] == 2
    assert row["실적 금액(백만)"] == 1.501
    assert result["extra_details"].empty


def test_midnight_and_empty_slots_are_kept(weekly_frame: pd.DataFrame) -> None:
    result = process_weekly(weekly_frame, "외장하드 주문.xlsx")["final"]
    assert result.query("`시작 시간` == '23:10'")["실적 수량"].sum() == 1
    assert (result["실적 수량"] == 0).any()


def test_weekly_conversion_uses_actual_quantity_and_views(weekly_frame: pd.DataFrame) -> None:
    stats = pd.DataFrame(
        [
            {
                "계정": "삼성공식파트너 쇼마젠시",
                "방송 ID": "100",
                "방송 제목": "외장하드 라이브",
                "방송일시": pd.Timestamp("2026-06-22 11:48"),
                "라이브중 시청수": 200,
                "유니크 결제자수": 5,
                "결제 상품수": 7,
                "데이터 업데이트 시각": pd.Timestamp("2026-06-23 10:00"),
            }
        ]
    )

    result = process_weekly(weekly_frame, "외장하드 주문.xlsx", live_stats=stats)
    matched = result["final"].query("`시작 시간` == '11:50'").iloc[0]
    unmatched = result["final"].query("`시작 시간` == '14:00'").iloc[0]

    assert matched["실적 View(만)"] == pytest.approx(0.02)
    assert matched["실적 전환율"] == pytest.approx(2 / 200)
    assert pd.isna(unmatched["실적 전환율"])
    assert result["live_stats"].query("`회차 시작 시간` == '11:50'").iloc[0]["실적 전환율(%)"] == pytest.approx(1.0)
    assert matched["목표 View(만)"] == pytest.approx(0.1)
    assert matched["목표 수량"] == pytest.approx(50)
    assert matched["목표 금액(백만)"] == pytest.approx(10)
    assert result["live_stats"].query("`회차 시작 시간` == '11:50'").iloc[0]["매칭 여부"] == "매칭"


def test_weekly_output_uses_target_and_actual_sections(weekly_frame: pd.DataFrame) -> None:
    final = process_weekly(weekly_frame, "외장하드 주문.xlsx")["final"]
    assert list(final.columns) == [
        "월", "일", "요일", "시작 시간", "duration", "ai 여부", "재방송여부", "채널", "방송주체",
        "운영그룹", "운영파트", "삼성 담당자", "거래선1", "거래선2", "품목1", "품목2", "비고",
        "목표 View(만)", "목표 수량", "목표 금액(백만)", "실적 View(만)", "실적 수량", "실적 전환율",
        "실적 금액(백만)", "비용률", "달성률", "제작(대행사)", "출연자1",
    ]
    manual_columns = [
        "재방송여부", "삼성 담당자", "거래선2", "비고",
    ]
    assert final[manual_columns].fillna("").eq("").all().all()
    assert set(final["목표 View(만)"]) == {0.1}
    assert set(final["목표 수량"]) == {50}
    assert set(final["목표 금액(백만)"]) == {10}
    assert set(final["ai 여부"]) == {"AI"}
    assert set(final["제작(대행사)"]) == {"쇼마젠시"}
    assert set(final["출연자1"]) == {"AI"}
    assert final["달성률"].eq(final["실적 금액(백만)"] / final["목표 금액(백만)"]).all()


@pytest.mark.parametrize("viewers", [None, 0, 200])
def test_weekly_conversion_zero_sales_and_missing_views(weekly_frame, viewers):
    stats = None if viewers is None else pd.DataFrame([{
        "방송일시": pd.Timestamp("2026-06-22 11:50"),
        "라이브중 시청수": viewers,
        "유니크 결제자수": 99,
        "데이터 업데이트 시각": pd.Timestamp("2026-06-23 10:00"),
    }])
    final = process_weekly(weekly_frame, "외장하드.xlsx", live_stats=stats)["final"]
    empty = final[(final["실적 수량"] == 0) & (final["실적 금액(백만)"] == 0)]
    assert not empty.empty
    assert empty["실적 전환율"].eq(0).all()
    populated = final.query("`시작 시간` == '11:50'").iloc[0]
    if viewers in (None, 0):
        assert pd.isna(populated["실적 전환율"])
    else:
        assert populated["실적 전환율"] == pytest.approx(2 / viewers)


def test_exact_duplicate_removed(weekly_frame: pd.DataFrame) -> None:
    duplicated = pd.concat([weekly_frame, weekly_frame.iloc[[0]]], ignore_index=True)
    result = process_weekly(duplicated, "외장하드.xlsx")
    assert len(result["duplicates"]) == 1
    assert result["final"]["실적 금액(백만)"].sum() == pytest.approx(1.502)


def test_shifted_live_detection(weekly_frame: pd.DataFrame) -> None:
    result = process_weekly(weekly_frame, "외장하드.xlsx")
    assert result["final"]["실적 수량"].sum() == 4


def test_multiple_raw_files_are_combined_and_deduplicated(weekly_frame: pd.DataFrame) -> None:
    first = weekly_frame.iloc[:3]
    second = weekly_frame.iloc[[0, 3, 4]]
    combined = pd.concat([first, second], ignore_index=True)
    result = process_weekly(combined, "외장하드 1.xlsx | 외장하드 2.xlsx")
    assert len(result["duplicates"]) == 1
    assert result["final"]["실적 수량"].sum() == 4


def test_mixed_weekly_file_types_are_rejected() -> None:
    with pytest.raises(ValueError, match="혼합"):
        infer_weekly_kind("외장하드.xlsx | 웨어러블.xlsx")


def test_period_file_name_infers_all_target_broadcast_dates() -> None:
    assert infer_target_dates("웨어러블 20260724~20260726 금~일 데이터 1.xlsx") == [
        pd.Timestamp("2026-07-24").date(),
        pd.Timestamp("2026-07-25").date(),
        pd.Timestamp("2026-07-26").date(),
    ]


def test_manual_slots_override_naver_download_timestamp_in_file_name() -> None:
    frame = pd.DataFrame(
        [["A", "2026-07-31 20:09", "상품A", 100_000, 0, "쇼핑라이브"]],
        columns=["주문번호", "결제일시", "상품명", "상품가격", "옵션가격", "주문 유입경로"],
    )
    custom_slots = {
        date(2026, 7, 31): (SlotRule("20:00", time(20, 0), time(21, 0)),),
    }

    result = process_weekly(
        frame,
        "스마트스토어_전체주문발주발송관리_20260803_0921.xlsx",
        "external",
        custom_slots=custom_slots,
    )

    assert set(map(tuple, result["final"][["월", "일"]].drop_duplicates().to_numpy())) == {(7, 31)}
    assert result["final"]["실적 수량"].sum() == 1
    assert result["excluded"].empty


def test_period_file_outputs_all_target_dates_and_excludes_after_last_midnight_window() -> None:
    frame = pd.DataFrame(
        [
            ["A", "2026-07-24 01:20", "상품A", 100_000, 0, "쇼핑라이브"],
            ["B", "2026-07-25 01:20", "상품B", 100_000, 0, "쇼핑라이브"],
            ["C", "2026-07-26 01:20", "상품C", 100_000, 0, "쇼핑라이브"],
            ["D", "2026-07-27 00:30", "상품D", 100_000, 0, "쇼핑라이브"],
            ["E", "2026-07-27 01:09", "상품E", 100_000, 0, "쇼핑라이브"],
        ],
        columns=["주문번호", "결제일시", "상품명", "상품가격", "옵션가격", "주문 유입경로"],
    )
    result = process_weekly(frame, "웨어러블 20260724~20260726 금~일 데이터 1.xlsx", "wearable")
    final = result["final"]
    assert set(map(tuple, final[["월", "일"]].drop_duplicates().to_numpy())) == {
        (7, 24), (7, 25), (7, 26),
    }
    assert final.groupby(["월", "일"])["실적 수량"].sum().to_dict() == {
        (7, 24): 1,
        (7, 25): 1,
        (7, 26): 2,
    }
    assert set(result["excluded"]["주문번호"]) == {"E"}


def test_product_order_number_duplicate_removed() -> None:
    frame = pd.DataFrame(
        [
            ["A", "P1", "2026-06-22 13:00", "상품1", "옵션", 1_000_000, 0, "쇼핑라이브"],
            ["A", "P1", "2026-06-22 13:00", "상품1", "옵션", 1_000_000, 0, "쇼핑라이브"],
        ],
        columns=["주문번호", "상품주문번호", "결제일시", "상품명", "옵션 정보", "상품가격", "옵션가격", "주문 유입경로"],
    )
    result = process_weekly(frame, "외장하드.xlsx")
    assert len(result["duplicates"]) == 1
    assert result["final"]["실적 금액(백만)"].sum() == pytest.approx(1)


def test_disabled_wearable_session_orders_are_excluded() -> None:
    frame = pd.DataFrame(
        [["W1", "2026-07-18 13:55", "웨어러블", 1, 100_000, 0, "쇼핑라이브"]],
        columns=["주문번호", "결제일시", "상품명", "수량", "상품가격", "옵션가격", "주문 유입경로"],
    )
    result = process_weekly(frame, "웨어러블.xlsx")
    assert result["final"].query("`시작 시간` == '13:50'")["실적 수량"].iloc[0] == 0
    assert len(result["excluded"]) == 1


def test_detail_filters_schedule_normalizes_orders_and_builds_three_outputs() -> None:
    frame = pd.DataFrame(
        [
            ["10,001", "2026-07-07 14:00", "워치9 40mm", "", "SM-L340NZEAKOO", 1_000_000, 0, "검색"],
            ["10-002", "2026-07-07 14:10", "워치9 40mm", "SM-L345NZKAKOO", "", 1_000_000, 0, "검색"],
            ["10/003", "2026-07-07 15:00", "워치9 44mm", "SM-L350NZKAKOO", "", 1_000_000, 0, "검색"],
            ["10.004", "2026-07-07 15:10", "워치9 44mm", "SM-L355NZKAKOO", "", 1_000_000, 0, "검색"],
            ["10 005", "2026-07-07 16:00", "울트라2", "SM-L715NZSAKOO", "", 1_000_000, 0, "검색"],
            ["20-001", "2026-07-07 16:30", "충전 어댑터", "EF-QF976CTEGKR", "", 500_000, 0, "검색"],
            ["30-001", "2026-07-08 14:00", "필터 밖 워치", "SM-L340NZEAKOO", "", 500_000, 0, "검색"],
            ["40-001", "2026-07-07 15:30", "미분류", "OPT-C", "SELLER-C", 500_000, 0, "쇼핑라이브"],
        ],
        columns=["주문번호", "결제일시", "상품명", "옵션관리코드", "판매자 상품코드", "상품가격", "옵션가격", "주문 유입경로"],
    )
    result = process_detail(
        frame,
        "주문.xlsx",
        date_range=(date(2026, 7, 7), date(2026, 7, 7)),
        time_range=(time(14, 0), time(17, 0)),
    )
    final = result["final"]
    assert list(final.columns) == ["버전", "결제일시", "주문번호", "상품명", "옵션 관리 코드", "금액"]
    assert set(result) >= {"basic", "wearable", "mobile_acc"}
    assert result["basic"].groupby("버전").size().to_dict() == {
        "워치9 40mm": 2,
        "워치9 44mm": 2,
        "울트라2": 1,
    }
    assert set(result["basic"]["주문번호"]) == {"10001", "10002", "10003", "10004", "10005"}
    assert set(result["wearable"]["주문번호"]) == {"10001", "10002", "10003", "10004", "10005"}
    assert set(result["mobile_acc"]["주문번호"]) == {"20001"}
    assert "버전" in result["basic"].columns
    assert "버전" not in result["wearable"].columns
    assert "버전" not in result["mobile_acc"].columns
    assert set(result["excluded"]["주문번호"]) == {"30-001", "40-001"}
    assert final["주문번호"].str.fullmatch(r"\d+").all()


def test_detail_uses_date_alias_and_short_sku_from_product_name() -> None:
    raw = pd.DataFrame(
        [
            {
                "주문번호": "10,001",
                "주문일시": "2026.07.30",
                "상품명": "삼성전자 갤럭시 워치9 블루투스 44mm SM-L350N 리뷰신세계2만+강화유리2매",
                "수량": 1,
                "판매금액": 350_000,
            }
        ]
    )
    frame = canonicalize_columns(raw, "주문 raw.xlsx")
    result = process_detail(frame, "주문 raw.xlsx")
    assert set(result["wearable"]["주문번호"]) == {"10001"}
    assert set(result["basic"]["버전"]) == {"워치9 44mm"}
    assert result["wearable"].loc[0, "옵션 관리 코드"] == "SM-L350N"


def test_weekly_optional_sku_filter_uses_product_name_code() -> None:
    frame = pd.DataFrame(
        [
            {
                "주문번호": "W1",
                "결제일시": "2026-07-23 01:20",
                "상품명": "갤럭시 워치9 44mm SM-L350N 사은품",
                "상품가격": 350_000,
                "옵션가격": 0,
                "주문 유입경로": "쇼핑라이브",
            }
        ]
    )
    result = process_weekly(
        frame,
        "웨어러블 20260723.xlsx",
        "wearable",
        allowed_skus={"SM-L350NZKAKOO"},
    )
    assert result["final"]["실적 수량"].sum() == 1


@pytest.mark.parametrize("kind,filename,expected", [
    ("external", "주문.xlsx", ("PP1", "NC", "Y3", "", 0.1, 50, 10)),
    ("wearable", "주문.xlsx", ("PP2", "모바일2", "갤럭시워치9", "갤럭시링", 0.2, 100, 24)),
    (None, "외장하드.xlsx", ("PP1", "NC", "Y3", "", 0.1, 50, 10)),
    (None, "웨어러블.xlsx", ("PP2", "모바일2", "갤럭시워치9", "갤럭시링", 0.2, 100, 24)),
])
def test_weekly_defaults_in_preview_and_download(weekly_frame, kind, filename, expected):
    from io import BytesIO
    from openpyxl import load_workbook
    from app import analyze_frame

    result, content, *_ = analyze_frame("weekly", kind, weekly_frame, filename, [])
    columns = ["운영그룹", "운영파트", "품목1", "품목2", "목표 View(만)", "목표 수량", "목표 금액(백만)"]
    defaults = dict(zip(columns, expected))
    defaults.update({"채널": "네이버", "방송주체": "SOP", "거래선1": "쇼마젠시", "제작(대행사)": "쇼마젠시", "출연자1": "AI"})
    final = result["final"]
    assert not final.empty
    for column, value in defaults.items():
        assert final[column].eq(value).all()
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        sheet = workbook["회차별 합계"]
        assert sheet.max_column == 39
        for row in range(8, sheet.max_row + 1):
            for column, value in defaults.items():
                from services.excel_writer import WEEKLY_EXCEL_COLUMNS
                assert sheet[f"{WEEKLY_EXCEL_COLUMNS[column]}{row}"].value == (None if value == "" else value)
    finally:
        workbook.close()
